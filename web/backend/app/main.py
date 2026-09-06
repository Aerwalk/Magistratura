import asyncio
import json
import logging
import os
import time
from datetime import datetime, timezone

from fastapi import FastAPI, WebSocket, WebSocketDisconnect, HTTPException, Query
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel

from config.constants import PRODUCTS, PRODUCT_COLORS, SHELF_CAPACITY, HOME_CELLS, SHELF_X, hex_to_rgba
from config.constants import BATTERY_CHARGE_RATE
from config.constants import add_product as add_product_const
import config.constants as _consts
from core.state import get_state, get_shelf_max_capacity, get_shelf_allowed_products
from core.grid import (
    get_grid_w, get_grid_h, get_exit_cell, get_inbound_cell,
    get_inbound_pick_cell, CUSTOM_SHELF_CELLS, REMOVED_SHELF_CELLS,
    all_shelves, set_exit_cell, set_inbound_cell, set_inbound_pick_cell,
    set_grid_size, add_custom_shelf, remove_custom_shelf, is_shelf_cell,
    rebuild_aisle_cells,
)
from core.pathfinding import astar
from entities.robot import robots, create_robot, Robot as RobotClass
from entities.uav import uav
from simulation.robot_controller import advance_robot_phases, _get_blocked_cells, _reset_tick_reservations
from simulation.task_manager import supply_labeled, supply_unlabeled, request_product, assign_tasks
from simulation.uav_controller import update_uav
from simulation.rebalancer import rebalance_warehouse, check_rebalance_complete
from agent_registry import agent_registry
from websocket_manager import ws_manager
from metrics import (record_cell_visit, record_distance_moved, record_charge_session,
                     record_battery_replaced, render_prometheus, record_sim_tick,
                     record_test_results)
import db
import otel
import obs

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

app = FastAPI(title="MagSim Warehouse")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

otel.init_tracing(app)
obs.init()

# Trace one simulation tick out of every N ticks so events hang off a
# `sim.tick` root span instead of flooding Tempo with per-tick roots.
SIM_TRACE_EVERY_TICKS = 200


class TelemetryBatch(BaseModel):
    robot_id: int
    samples: list = []


class LogEvent(BaseModel):
    robot_id: int | None = None
    level: str = "info"
    event: str
    payload: dict | None = None


def _parse_ts(value):
    if not value:
        return None
    try:
        text = str(value).replace("Z", "+00:00")
        return datetime.fromisoformat(text)
    except Exception:
        return None


def color_to_hex(rgba):
    r, g, b, a = rgba
    return '#{:02x}{:02x}{:02x}'.format(int(r*255), int(g*255), int(b*255))


def serialize_state():
    state = get_state()
    robot_data = []
    for r in robots:
        robot_data.append({
            "id": r.id,
            "name": r.name,
            "product": r.product,
            "products": list(r.products),
            "color": color_to_hex(r.color),
            "pos": list(r.pos),
            "draw_pos": r.draw_pos,
            "phase": r.phase,
            "carry": r.carry,
            "carry_qty": r.carry_qty,
            "home": list(r.home),
            "speed": r.speed,
            "capacity": r.capacity,
            "battery": r.battery,
            "battery_max": r.battery_max,
            "charging": r.phase == "charging",
            "stuck": r.stuck_alert_raised,
            "stationary_secs": round(r.stationary_ticks * 0.05, 1),
            "distance_moved": round(r.distance_moved, 1),
            "charge_eta_sec": (
                round(max(0.0, (r.battery_max - r.battery) / BATTERY_CHARGE_RATE), 1)
                if r.phase == "charging" else None
            ),
            "task": str(r.task) if r.task else None,
            "target_shelf": r.target_shelf,
            "carry_bob": r.carry_bob,
            "carry_appear_scale": r.carry_appear_scale,
            "delivery_anim_ticks": r.delivery_anim_ticks,
            "pickup_anim_ticks": r.pickup_anim_ticks,
        })
    shelves = []
    for shelf_pos in all_shelves():
        info = state.shelf_state.get(shelf_pos, {"product": None, "count": 0})
        allowed = get_shelf_allowed_products(shelf_pos)
        shelves.append({
            "x": shelf_pos[0], "y": shelf_pos[1],
            "product": info["product"],
            "count": info["count"],
            "capacity": get_shelf_max_capacity(shelf_pos),
            "allowed_products": list(allowed) if allowed else None,
        })
    products = []
    for p in PRODUCTS:
        products.append({
            "name": p,
            "color": color_to_hex(PRODUCT_COLORS.get(p, (0.5, 0.5, 0.5, 1.0))),
            "stock_total": state.stock_total.get(p, 0),
            "demand": state.demand.get(p, 0),
            "delivered": state.delivered_total.get(p, 0),
            "queue_length": len(state.robot_task_queues.get(p, [])),
            "retrieved": state.retrieved_total.get(p, 0),
            "on_shelves": sum(
                1 for info in state.shelf_state.values()
                if info["product"] == p
            ),
        })
    return {
        "type": "state_update",
        "data": {
            "robots": robot_data,
            "uav": {
                "pos": list(uav.pos),
                "state": uav.state,
                "scan_ticks": uav.scan_ticks,
                "uav_color": uav.uav_color,
            },
            "shelves": shelves,
            "grid_info": {
                "width": get_grid_w(),
                "height": get_grid_h(),
                "exit_cell": list(get_exit_cell()),
                "inbound_cell": list(get_inbound_cell()),
                "inbound_pick_cell": list(get_inbound_pick_cell()),
                "uav_home": [7, 5, 4.5],
                "shelf_x": sorted(SHELF_X),
                "home_cells": [list(c) for c in HOME_CELLS],
                "custom_shelves": [list(c) for c in CUSTOM_SHELF_CELLS],
                "removed_shelves": [list(c) for c in REMOVED_SHELF_CELLS],
            },
            "stats": {
                "stock_total": dict(state.stock_total),
                "demand": dict(state.demand),
                "delivered_total": dict(state.delivered_total),
                "retrieved_total": dict(state.retrieved_total),
                "rebalance_active": state.rebalance_active,
                "rebalance_moves_remaining": state.rebalance_moves_remaining,
                "shelf_capacity": SHELF_CAPACITY,
                "stuck_robots": sum(1 for rb in robots if rb.stuck_alert_raised),
                "battery_replaced_total": sum(BATTERY_REPLACED_MEM.values()),
            },
            "products": products,
            "log": list(state.event_log),
        }
    }


_TICK = 0
TELEMETRY_EVERY_TICKS = 20  # emit ~1 Hz at 50 ms ticks
_START_TIME = time.time()

# Automated tests: run in an isolated subprocess and surface as Prometheus gauges.
AUTO_TESTS_INTERVAL_SEC = float(os.getenv("AUTO_TESTS_INTERVAL_SEC", "600"))
AUTO_TESTS_FIRST_DELAY_SEC = float(os.getenv("AUTO_TESTS_FIRST_DELAY_SEC", "60"))
_TESTS_CWD = os.path.dirname(os.path.abspath(__file__))

# Live count of battery replacements (fallback when DB is off / unflushed).
# DB robot_logs 'battery_replaced' events remain the persisted source.
BATTERY_REPLACED_MEM: dict[int, int] = {}


def emit_telemetry() -> None:
    ts = datetime.now(timezone.utc)
    for r in robots:
        db.push_telemetry(
            ts, r.id,
            float(r.pos[0]), float(r.pos[1]),
            r.phase, r.battery, r.battery_max,
            r.phase == "charging",
            r.speed,
            r.task["type"] if r.task else None,
        )


def _find_robot(rid: int):
    for r in robots:
        if r.id == rid:
            return r
    return None


def _apply_uav_scan(qty: int, product: str) -> None:
    state = get_state()
    state.event_log.appendleft(
        f"БВС распознал товар без метки как {product}, количество {qty}"
    )
    has_robot = any(product in r.products for r in robots)
    if not has_robot:
        state.event_log.appendleft(f"⚠ ВНИМАНИЕ: нет РТС, способного возить {product}!")
    for _ in range(qty):
        state.stock_total[product] += 1
        state.robot_task_queues[product].append(
            {"type": "putaway", "product": product, "source": "unlabeled"}
        )


async def handle_agent_message(kind: str, msg: dict) -> None:
    """Dispatch an agent message; wraps non-telemetry handling in an OTel span
    that nests under the sender's own span (via W3C traceparent)."""
    span_handle = None
    if msg.get("type") != "telemetry":
        span_handle = obs.begin_span(
            f"agent.msg.{msg.get('type')}",
            kind=kind,
            robot_id=msg.get("robot_id"),
            parent_ctx=obs.extract_traceparent(msg),
        )
    try:
        await _handle_agent_message_inner(kind, msg)
    except Exception:
        obs.end_span(span_handle)
        raise
    obs.end_span(span_handle)


async def _handle_agent_message_inner(kind: str, msg: dict) -> None:
    mtype = msg.get("type")
    if kind == "robot":
        r = _find_robot(msg.get("robot_id"))
        if r is None:
            return
        if mtype == "ready":
            r.agent_ready = True
            r.route = []
            r.current_target = None
            r.executing = False
            r.wait_ticks = 0
            if r.phase == "charging":
                r.phase = "idle"
                r.charging_ticks = 0
                r.charging_start_ts = None
            get_state().event_log.appendleft(f"РТС-{r.id}: агент подключён")
        elif mtype == "telemetry":
            prev_draw = r.draw_pos
            new_draw = [float(msg["pos"][0]), float(msg["pos"][1])]
            dx = abs(new_draw[0] - prev_draw[0]) + abs(new_draw[1] - prev_draw[1])
            r.draw_pos = new_draw
            new_pos = (int(round(new_draw[0])), int(round(new_draw[1])))
            if new_pos != r.pos:
                record_cell_visit(r.id, new_pos)
            if dx > 0:
                r.distance_moved += dx
                record_distance_moved(r.id, dx)
            if isinstance(msg.get("speed"), (int, float)):
                r.speed = max(0.0, float(msg["speed"]))
            r.pos = new_pos
            r.battery = float(msg["battery"])
            r.battery_max = float(msg["battery_max"])
        elif mtype == "charging_progress":
            payload = msg.get("payload") or {}
            db.push_log(r.id, "info", "charging_progress", {
                "battery": round(payload.get("battery", r.battery), 2),
                "battery_max": round(payload.get("battery_max", r.battery_max), 2),
                "progress_pct": payload.get("progress_pct"),
                "remaining_sec": payload.get("remaining_sec"),
                "elapsed_sec": payload.get("elapsed_sec"),
            })
        elif mtype == "arrived":
            pos = msg["pos"]
            r.pos = (int(round(pos[0])), int(round(pos[1])))
            r.draw_pos = [float(pos[0]), float(pos[1])]
            r.route = []
            r.current_target = None
            r.executing = False
            r.wait_ticks = 2
            r.stuck_ticks = 0
            record_cell_visit(r.id, r.pos)
        elif mtype == "charge_end":
            payload = msg["payload"]
            prev_max = r.battery_max
            r.battery = float(payload["battery_to"])
            r.battery_max = float(payload["battery_max_after"])
            r.phase = "idle"
            r.charging_ticks = 0
            start_ts = r.charging_start_ts
            r.charging_start_ts = None
            r.battery_at_charge_start = 0.0
            record_charge_session(r)
            if r.battery_max < prev_max - 0.01:
                record_battery_replaced(r)
            db.push_log(r.id, "info", "charging_end", payload)
            db.push_charge_session(r.id, start_ts, datetime.now(timezone.utc),
                                   payload["duration_sec"], payload["battery_from"],
                                   payload["battery_to"], payload["battery_max_after"])
            get_state().event_log.appendleft(
                f"РТС-{r.id} заряжен до {payload['battery_to']}% за {payload['duration_sec']:.0f}с "
                f"(макс. ёмкость {payload['battery_max_after']:.1f}%)"
            )
    elif kind == "uav":
        if mtype == "telemetry":
            uav.pos = [float(x) for x in msg["pos"]]
            uav.state = msg.get("state", uav.state)
            uav.scan_ticks = int(msg.get("scan_ticks", 0))
        elif mtype == "scan_done":
            _apply_uav_scan(msg.get("qty", 0), msg.get("product"))


async def _assign_pending_robots() -> None:
    while agent_registry.has_spare():
        r = next((x for x in robots if x.agent_ws is None), None)
        if r is None:
            break
        ws = agent_registry.pop_spare()
        r.agent_ws = ws
        agent_registry.assign(r.id, ws)
        with obs.span("agent.assign", robot_id=r.id):
            ok = await agent_registry.send_robot(r.id, {
                "cmd": "assign",
                "robot_id": r.id,
                "home": list(r.home),
                "speed": r.speed,
            })
        if not ok:
            r.agent_ws = None
            r.agent_ready = False


async def agent_bootstrapper():
    while True:
        try:
            await _assign_pending_robots()
        except Exception as e:
            logger.error(f"Agent bootstrap error: {e}")
        await asyncio.sleep(1.0)


async def flush_routes() -> None:
    for r in robots:
        if not r.agent_ready or r.executing or r.phase == "charging" or not r.route:
            continue
        r.executing = True
        r.current_target = r.route[0]
        ok = await agent_registry.send_robot(r.id, {
            "cmd": "route",
            "route": [list(p) for p in r.route],
        })
        if not ok:
            r.executing = False
            r.current_target = None


async def simulation_tick():
    global _TICK
    while True:
        try:
            state = get_state()
            _reset_tick_reservations()
            with obs.span("sim.tick", tick=_TICK) if (_TICK % SIM_TRACE_EVERY_TICKS == 0) else obs.null():
                advance_robot_phases()
                assign_tasks()
                update_uav()
                check_rebalance_complete()
            await flush_routes()
            _TICK += 1
            record_sim_tick(time.time())
            if _TICK % TELEMETRY_EVERY_TICKS == 0:
                emit_telemetry()
            payload = serialize_state()
            await ws_manager.broadcast(payload)
        except Exception as e:
            logger.error(f"Simulation tick error: {e}")
        await asyncio.sleep(0.05)


@app.on_event("startup")
async def startup():
    global _START_TIME
    _START_TIME = time.time()
    asyncio.create_task(db.init(forever=True))
    asyncio.create_task(db.writer_loop())
    asyncio.create_task(db.loki_writer_loop())
    asyncio.create_task(simulation_tick())
    asyncio.create_task(agent_bootstrapper())
    asyncio.create_task(auto_test_runner())


@app.on_event("shutdown")
async def shutdown():
    await db.close()


@app.websocket("/ws")
async def websocket_endpoint(websocket: WebSocket):
    await ws_manager.connect(websocket)
    try:
        while True:
            data = await websocket.receive_text()
            msg = json.loads(data)
            msg_type = msg.get("type", "")
            msg_data = msg.get("data", {})
            try:
                state = get_state()
                ui_span = obs.begin_span(f"ui.{msg_type}", user_action=msg_type)
                if msg_type == "supply_labeled":
                    supply_labeled(msg_data["product"], msg_data["qty"])
                elif msg_type == "supply_unlabeled":
                    supply_unlabeled(msg_data["qty"])
                elif msg_type == "request_product":
                    request_product(msg_data["product"], msg_data["qty"])
                elif msg_type == "create_robot":
                    robot = create_robot(msg_data["product"])
                    if robot is None:
                        get_state().event_log.appendleft("⚠ Нет свободной домашней станции для нового РТС!")
                    elif not agent_registry.has_spare():
                        get_state().event_log.appendleft(f"⚠ Нет свободного агента для РТС-{robot.id}!")
                    else:
                        ws = agent_registry.pop_spare()
                        robot.agent_ws = ws
                        agent_registry.assign(robot.id, ws)
                        await agent_registry.send_robot(robot.id, {
                            "cmd": "assign",
                            "robot_id": robot.id,
                            "home": list(robot.home),
                            "speed": robot.speed,
                        })
                elif msg_type == "remove_robot":
                    rid = msg_data["robot_id"]
                    target = next((r for r in robots if r.id == rid), None)
                    if target is None:
                        get_state().event_log.appendleft(f"РТС-{rid} не найден, нечего удалять")
                    elif len(robots) <= 1:
                        get_state().event_log.appendleft("⚠ Нельзя удалить последнего РТС!")
                    else:
                        # Return cargo/task to queues (sync mutations first)
                        if target.carry:
                            state.stock_total[target.carry] += 1
                            state.robot_task_queues[target.carry].append(
                                {"type": "putaway", "product": target.carry}
                            )
                        task = target.task
                        if task and task.get("type") in ("putaway", "retrieve"):
                            state.robot_task_queues[task["product"]].append(task)
                        name = target.name
                        released_ws = agent_registry.release_robot(rid)
                        robots.remove(target)
                        from metrics import clear_robot
                        clear_robot(rid)
                        get_state().event_log.appendleft(
                            f"РТС-{rid} ({name}) выведен из парка, задачи возвращены в очередь"
                        )
                        # Tell the released agent to stand by (await after mutations)
                        if released_ws is not None:
                            try:
                                await released_ws.send_text(json.dumps({"cmd": "stop"}))
                            except Exception:
                                pass
                elif msg_type == "create_product":
                    add_product_const(msg_data["name"], msg_data["color"])
                elif msg_type == "rebalance":
                    rebalance_warehouse(force=True)
                elif msg_type == "reset_robot":
                    rid = msg_data["robot_id"]
                    for r in robots:
                        if r.id == rid:
                            carry = r.carry
                            task = r.task
                            # Return cargo to stock + create putaway task
                            if carry:
                                state.stock_total[carry] += 1
                                state.robot_task_queues[carry].appendleft({"type": "putaway", "product": carry})
                                state.event_log.appendleft(f"РТС-{rid} сброшен, груз {carry} возвращён на станцию забора")
                            # Return task to queue
                            if task and task.get("type") == "putaway":
                                p = task["product"]
                                state.robot_task_queues[p].appendleft(task)
                            elif task and task.get("type") == "retrieve":
                                p = task["product"]
                                state.robot_task_queues[p].appendleft(task)
                            # Clear robot state without teleporting home
                            r.task = None
                            r.carry = None
                            r.target_shelf = None
                            r.route = []
                            r.current_target = None
                            r.phase = "idle"
                            r.delivery_anim_ticks = 0
                            r.pickup_anim_ticks = 0
                            r.carry_appear_scale = 1.0
                            r.wait_ticks = 0
                            r.stuck_ticks = 0
                            r.stationary_ticks = 0
                            r.stuck_alert_raised = False
                            r.stuck_since = None
                            r.executing = False
                            await agent_registry.send_robot(rid, {"cmd": "stop"})
                            # Send robot home
                            if r.pos != r.home:
                                blocked = _get_blocked_cells(exclude_robot=r)
                                route = astar(r.pos, r.home, blocked=blocked)
                                if route:
                                    r.route = route
                                    r.phase = "to_home"
                            if not carry and not task:
                                state.event_log.appendleft(f"РТС-{rid} сброшен, возвращается на станцию")
                            break
                elif msg_type == "replace_battery":
                    rid = msg_data["robot_id"]
                    for r in robots:
                        if r.id == rid:
                            old_max = r.battery_max
                            r.battery = 100.0
                            r.battery_max = 100.0
                            r.charging_ticks = 0
                            r.charging_start_ts = None
                            r.battery_at_charge_start = 0.0
                            if r.phase == "charging":
                                r.phase = "idle"
                            r.stationary_ticks = 0
                            r.stuck_alert_raised = False
                            r.stuck_since = None
                            BATTERY_REPLACED_MEM[rid] = BATTERY_REPLACED_MEM.get(rid, 0) + 1
                            record_battery_replaced(r)
                            get_state().event_log.appendleft(
                                f"АКБ РТС-{rid} заменён: ёмкость {old_max:.1f}% → 100%"
                            )
                            db.push_log(rid, "info", "battery_replaced",
                                        {"old_battery_max": round(old_max, 1), "new_battery_max": 100.0})
                            await agent_registry.send_robot(rid, {"cmd": "replace_battery"})
                            break
                elif msg_type == "set_robot_product":
                    rid = msg_data["robot_id"]
                    product = msg_data["product"]
                    enabled = msg_data["enabled"]
                    for r in robots:
                        if r.id == rid:
                            if enabled:
                                r.products.add(product)
                            else:
                                r.products.discard(product)
                            # sync primary product (None if empty)
                            if r.products:
                                first = next(iter(r.products))
                                r.product = first
                                if first in PRODUCT_COLORS:
                                    r.color = PRODUCT_COLORS[first]
                            else:
                                r.product = None
                            break
                elif msg_type == "set_robot_color":
                    rid = msg_data["robot_id"]
                    color_hex = msg_data.get("color", "")
                    for r in robots:
                        if r.id == rid:
                            if color_hex and color_hex.lower() != "#ffffff":
                                from config.constants import hex_to_rgba
                                r.color = hex_to_rgba(color_hex)
                            else:
                                # reset to product color
                                if r.product and r.product in PRODUCT_COLORS:
                                    r.color = PRODUCT_COLORS[r.product]
                            break
                elif msg_type == "set_robot_speed":
                    rid = msg_data["robot_id"]
                    speed = max(0, min(1.0, msg_data["speed"]))
                    for r in robots:
                        if r.id == rid:
                            r.speed = speed
                            await agent_registry.send_robot(rid, {"cmd": "set_speed", "speed": r.speed})
                            break
                elif msg_type == "set_robot_capacity":
                    rid = msg_data["robot_id"]
                    cap = max(1, msg_data["capacity"])
                    for r in robots:
                        if r.id == rid:
                            r.capacity = cap
                            break
                elif msg_type == "set_grid_size":
                    set_grid_size(msg_data["width"], msg_data["height"])
                elif msg_type == "set_exit_cell":
                    set_exit_cell(msg_data["x"], msg_data["y"])
                elif msg_type == "set_inbound_cell":
                    set_inbound_cell(msg_data["x"], msg_data["y"])
                elif msg_type == "set_inbound_pick_cell":
                    set_inbound_pick_cell(msg_data["x"], msg_data["y"])
                elif msg_type == "add_custom_shelf":
                    add_custom_shelf(msg_data["x"], msg_data["y"])
                elif msg_type == "remove_custom_shelf":
                    remove_custom_shelf(msg_data["x"], msg_data["y"])
                elif msg_type == "rename_robot":
                    rid = msg_data["robot_id"]
                    new_name = msg_data["name"]
                    for r in robots:
                        if r.id == rid:
                            r.name = new_name
                            get_state().event_log.appendleft(f"РТС-{rid} переименован в {new_name}")
                            break
                elif msg_type == "set_shelf_capacity":
                    x = int(msg_data["x"])
                    y = int(msg_data["y"])
                    cap = int(msg_data["capacity"])
                    if (x, y) in state.shelf_state:
                        if cap <= 0:
                            state.shelf_capacity_override.pop((x, y), None)
                        else:
                            state.shelf_capacity_override[(x, y)] = cap
                        state.event_log.appendleft(f"Вместимость полки ({x},{y}) = {cap}")
                elif msg_type == "set_shelf_allowed_products":
                    x = int(msg_data["x"])
                    y = int(msg_data["y"])
                    products_list = msg_data.get("products")
                    if products_list is None:
                        state.shelf_allowed_products.pop((x, y), None)
                    else:
                        state.shelf_allowed_products[(x, y)] = set(products_list)
                    state.event_log.appendleft(f"Ограничения полки ({x},{y}) обновлены")
                elif msg_type == "set_uav_color":
                    color_hex = msg_data.get("color")
                    if color_hex and color_hex.lower() != "#ffffff":
                        uav.uav_color = color_hex
                    else:
                        uav.uav_color = None
                    get_state().event_log.appendleft(f"Цвет БВС изменён")
                elif msg_type == "set_robot_home":
                    rid = msg_data["robot_id"]
                    hx = int(msg_data["home_x"])
                    hy = int(msg_data["home_y"])
                    for r in robots:
                        if r.id == rid:
                            old = r.home
                            r.home = (hx, hy)
                            if old in _consts.HOME_CELLS:
                                _consts.HOME_CELLS.remove(old)
                            if (hx, hy) not in _consts.HOME_CELLS:
                                _consts.HOME_CELLS.append((hx, hy))
                            from core.grid import rebuild_aisle_cells
                            rebuild_aisle_cells()
                            get_state().event_log.appendleft(f"РТС-{rid} новая станция: ({hx},{hy})")
                            await agent_registry.send_robot(rid, {"cmd": "set_home", "home": [hx, hy]})
                            break
                obs.end_span(ui_span)
            except Exception as e:
                obs.end_span(ui_span)
                logger.error(f"Handler error for {msg_type}: {e}")
                get_state().event_log.appendleft(f"⚠ Ошибка: {e}")
    except WebSocketDisconnect:
        ws_manager.disconnect(websocket)
    except Exception as e:
        logger.error(f"WS error: {e}")
        ws_manager.disconnect(websocket)


@app.websocket("/agent/ws/{kind}")
async def agent_ws_endpoint(kind: str, websocket: WebSocket):
    await websocket.accept()
    if kind == "robot":
        agent_registry.add_spare(websocket)
    elif kind == "uav":
        agent_registry.uav_ws = websocket
    else:
        await websocket.close(code=1008)
        return
    try:
        while True:
            data = await websocket.receive_text()
            msg = json.loads(data)
            try:
                await handle_agent_message(kind, msg)
            except Exception as e:
                logger.error(f"Agent handler error ({kind}): {e}")
    except WebSocketDisconnect:
        pass
    except Exception as e:
        logger.error(f"Agent WS error ({kind}): {e}")
    finally:
        agent_registry.remove_ws(websocket)
        for r in robots:
            if r.agent_ws == websocket:
                r.agent_ws = None
                r.agent_ready = False
                r.executing = False
                r.current_target = None


@app.get("/api/stats")
def get_full_stats():
    state = get_state()
    products_detail = []
    for p in PRODUCTS:
        on_shelf_count = sum(
            info["count"] for info in state.shelf_state.values()
            if info["product"] == p
        )
        products_detail.append({
            "name": p,
            "color": color_to_hex(PRODUCT_COLORS.get(p, (0.5, 0.5, 0.5, 1.0))),
            "stock_total": state.stock_total.get(p, 0),
            "demand": state.demand.get(p, 0),
            "delivered": state.delivered_total.get(p, 0),
            "retrieved": state.retrieved_total.get(p, 0),
            "queue_length": len(state.robot_task_queues.get(p, [])),
            "on_shelves": on_shelf_count,
        })
    robots_detail = []
    for r in robots:
        robots_detail.append({
            "id": r.id,
            "name": r.name,
            "product": r.product,
            "products": list(r.products),
            "color": color_to_hex(r.color),
            "phase": r.phase,
            "carry": r.carry,
            "carry_qty": r.carry_qty,
            "speed": r.speed,
            "capacity": r.capacity,
            "pos": list(r.pos),
            "home": list(r.home),
        })
    shelves_detail = []
    for shelf_pos in all_shelves():
        info = state.shelf_state.get(shelf_pos, {"product": None, "count": 0})
        allowed = get_shelf_allowed_products(shelf_pos)
        shelves_detail.append({
            "pos": list(shelf_pos),
            "x": shelf_pos[0], "y": shelf_pos[1],
            "product": info["product"],
            "count": info["count"],
            "capacity": get_shelf_max_capacity(shelf_pos),
            "allowed_products": list(allowed) if allowed else None,
        })
    return {
        "products": products_detail,
        "robots": robots_detail,
        "shelves": shelves_detail,
        "rebalance_active": state.rebalance_active,
        "rebalance_moves_remaining": state.rebalance_moves_remaining,
    }


@app.get("/api/export/excel")
def export_excel():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.chart import BarChart, Reference
    from io import BytesIO
    import datetime

    state = get_state()
    wb = Workbook()

    # ─── Sheet 1: Products ───
    ws1 = wb.active
    ws1.title = "Материалы"
    headers = ["Материал", "Всего поставлено", "Спрос", "Доставлено на склад",
               "Забрано со склада", "В очереди", "На полках"]
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2c3e50", end_color="2c3e50", fill_type="solid")
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )

    for col, h in enumerate(headers, 1):
        cell = ws1.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')
        cell.border = thin_border

    for row_idx, p in enumerate(PRODUCTS, 2):
        vals = [
            p,
            state.stock_total.get(p, 0),
            state.demand.get(p, 0),
            state.delivered_total.get(p, 0),
            state.retrieved_total.get(p, 0),
            len(state.robot_task_queues.get(p, [])),
            sum(1 for info in state.shelf_state.values() if info["product"] == p),
        ]
        for col, v in enumerate(vals, 1):
            cell = ws1.cell(row=row_idx, column=col, value=v)
            cell.border = thin_border
            cell.alignment = Alignment(horizontal='center')

    # Bar chart for product stats
    chart = BarChart()
    chart.type = "col"
    chart.title = "Сводка по материалам"
    chart.y_axis.title = "Количество"
    chart.x_axis.title = "Материалы"
    data = Reference(ws1, min_col=2, max_col=7, min_row=1, max_row=len(PRODUCTS)+1)
    cats = Reference(ws1, min_col=1, min_row=2, max_row=len(PRODUCTS)+1)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.shape = 4
    ws1.add_chart(chart, "A10")

    # ─── Sheet 2: Robots ───
    ws2 = wb.create_sheet("РТС")
    r_headers = ["ID", "Имя", "Продукт", "Продукты", "Фаза", "Несёт",
                 "Скорость", "Вместимость", "Позиция", "Дом"]
    for col, h in enumerate(r_headers, 1):
        cell = ws2.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')
        cell.border = thin_border
    for row_idx, r in enumerate(robots, 2):
        vals = [
            r.id, r.name, r.product, ", ".join(r.products), r.phase,
            r.carry or "", r.speed, r.capacity,
            f"({r.pos[0]}, {r.pos[1]})",
            f"({r.home[0]}, {r.home[1]})",
        ]
        for col, v in enumerate(vals, 1):
            cell = ws2.cell(row=row_idx, column=col, value=v)
            cell.border = thin_border
            cell.alignment = Alignment(horizontal='center')

    # ─── Sheet 3: Shelves ───
    ws3 = wb.create_sheet("Полки")
    s_headers = ["X", "Y", "Продукт", "Количество"]
    for col, h in enumerate(s_headers, 1):
        cell = ws3.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')
        cell.border = thin_border
    row_idx = 2
    for shelf_pos in all_shelves():
        info = state.shelf_state.get(shelf_pos, {"product": None, "count": 0})
        vals = [shelf_pos[0], shelf_pos[1], info["product"] or "", info["count"]]
        for col, v in enumerate(vals, 1):
            cell = ws3.cell(row=row_idx, column=col, value=v)
            cell.border = thin_border
            cell.alignment = Alignment(horizontal='center')
        row_idx += 1

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    from fastapi.responses import Response
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f"attachment; filename=magsim_report_{datetime.date.today().isoformat()}.xlsx"
        },
    )


import subprocess
import sys


def _run_tests_subprocess(mode: str, scenario_id: int | None = None) -> dict:
    """Run tests/scenarios in an isolated subprocess so the live sim is untouched."""
    if mode == "all":
        code = (
            "import json; "
            "from tests.runner import run_everything; "
            "print(json.dumps(run_everything(), ensure_ascii=False))"
        )
    else:
        code = (
            "import json; "
            "from tests.scenarios import run_scenario; "
            f"print(json.dumps(run_scenario({scenario_id}), ensure_ascii=False))"
        )
    proc = None
    try:
        proc = subprocess.run(
            [sys.executable, "-c", code],
            cwd=_TESTS_CWD,
            capture_output=True,
            text=True,
            timeout=300,
        )
        out = (proc.stdout or "").strip()
        start = out.find("{")
        if start >= 0:
            payload = json.loads(out[start:])
        else:
            payload = {}
    except Exception as e:
        payload = {}
    if not payload:
        err = ""
        if proc is not None:
            err = (proc.stderr or proc.stdout or "")[-2000:]
        payload = {"error": (err or "subprocess returned no JSON")}
    return payload


async def auto_test_runner():
    """Periodically runs the full test suite in a subprocess and records results
    into Prometheus gauges (magsim_tests_*) so they are visible in Grafana."""
    await asyncio.sleep(AUTO_TESTS_FIRST_DELAY_SEC)
    while True:
        try:
            payload = await asyncio.to_thread(_run_tests_subprocess, "all")
            record_test_results(payload)
            tests = payload.get("tests") or {}
            scenarios = payload.get("scenarios") or []
            failed_sc = sum(1 for s in scenarios if not s.get("passed"))
            logger.info(
                f"Auto tests: unit {tests.get('passed', 0)}/{tests.get('total', 0)} "
                f"passed, scenarios {len(scenarios) - failed_sc}/{len(scenarios)} passed"
            )
        except Exception as e:
            logger.error(f"Auto tests error: {e}")
        await asyncio.sleep(AUTO_TESTS_INTERVAL_SEC)


@app.get("/api/run-tests")
def run_all_tests():
    """Run all unittest tests and scenarios in an isolated subprocess."""
    return _run_tests_subprocess("all")


@app.get("/api/scenarios")
def list_scenarios():
    """Return list of available scenarios."""
    from tests.scenarios import SCENARIO_META
    return [
        {"id": sid, "name": meta["name"], "expected": meta["expected"]}
        for sid, meta in sorted(SCENARIO_META.items())
    ]


@app.get("/api/run-scenario/{scenario_id}")
def run_one_scenario(scenario_id: int):
    """Run a single scenario in an isolated subprocess by ID."""
    result = _run_tests_subprocess("scenario", scenario_id)
    if result.get("error"):
        raise HTTPException(status_code=500, detail=result["error"])
    if result.get("id") != scenario_id:
        raise HTTPException(status_code=404, detail=f"Scenario {scenario_id} not found")
    return result


@app.get("/api/products")
def get_products():
    return PRODUCTS


@app.post("/api/telemetry")
async def ingest_telemetry(batch: TelemetryBatch):
    now = datetime.now(timezone.utc)
    queued = 0
    for s in batch.samples[:500]:
        ts = _parse_ts(s.get("ts")) if isinstance(s, dict) else None
        ts = ts or now
        db.push_telemetry(
            ts, batch.robot_id,
            float(s.get("x", 0.0)), float(s.get("y", 0.0)),
            s.get("phase"), s.get("battery"), s.get("battery_max"),
            bool(s.get("charging", False)),
            float(s.get("speed", 0.0)), s.get("task_type"),
        )
        queued += 1
    return {"ok": True, "queued": queued}


@app.post("/api/log")
async def ingest_log(ev: LogEvent):
    db.push_log(ev.robot_id, ev.level, ev.event, ev.payload)
    return {"ok": True}


@app.get("/api/health")
async def health():
    return {"status": "ok", "db_enabled": db.is_enabled(), "robots": len(robots)}


@app.get("/healthz")
async def healthz():
    connected = sum(1 for r in robots if getattr(r, "agent_ws", None) is not None)
    ready = sum(1 for r in robots if getattr(r, "agent_ready", False))
    return {
        "status": "ok" if (connected == len(robots) and db.is_enabled()) else "degraded",
        "db_enabled": db.is_enabled(),
        "agents_connected": connected,
        "robots": len(robots),
        "robots_ready": ready,
        "ticks": _TICK,
        "uptime_sec": round(time.time() - _START_TIME),
        "time": datetime.now(timezone.utc).isoformat(),
    }


@app.get("/metrics")
async def metrics():
    from fastapi.responses import Response
    body, content_type = render_prometheus()
    return Response(content=body, media_type=content_type)


@app.get("/api/debug")
async def debug_state():
    from agent_registry import agent_registry
    robots_detail = []
    for r in robots:
        robots_detail.append({
            "id": r.id,
            "name": r.name,
            "phase": r.phase,
            "pos": list(r.pos),
            "home": list(r.home),
            "task": r.task,
            "route": [list(p) for p in r.route],
            "current_target": list(r.current_target) if r.current_target else None,
            "executing": r.executing,
            "wait_ticks": r.wait_ticks,
            "stuck_ticks": r.stuck_ticks,
            "battery": round(r.battery, 2),
            "battery_max": round(r.battery_max, 2),
            "agent_ready": r.agent_ready,
            "agent_ws": r.agent_ws is not None,
            "carry": r.carry,
            "target_shelf": r.target_shelf,
        })
    state = get_state()
    return {
        "robots": robots_detail,
        "spare_ws": len(agent_registry.spare_robot_ws),
        "assigned_ws": len(agent_registry.robot_agents),
        "queues": {p: len(q) for p, q in state.robot_task_queues.items()},
        "uav_scan_queue": list(state.uav_scan_queue),
        "event_log": list(state.event_log)[:20],
    }


@app.get("/api/predictions")
async def predictions(window_hours: int = 1):
    return await db.get_predictions(window_hours)


@app.get("/api/dashboard")
async def dashboard_api():
    """Aggregated data for the frontend dashboard page (charts, warnings, counters)."""
    from config.constants import BATTERY_DEGRADED_WARNING_PCT, BATTERY_LOW

    history = await db.get_battery_history(hours=3, bucket="1 minute")
    db_counts = await db.count_battery_replacements() if db.is_enabled() else {}
    replaced: dict[int, int] = {}
    for rid in set(db_counts) | set(BATTERY_REPLACED_MEM):
        replaced[rid] = max(db_counts.get(rid, 0), BATTERY_REPLACED_MEM.get(rid, 0))

    robots_detail = []
    degraded_count = 0
    critical_count = 0
    for r in robots:
        degradation = round(100.0 - r.battery_max, 1)
        degraded = degradation > BATTERY_DEGRADED_WARNING_PCT
        low = r.battery <= BATTERY_LOW
        health = "critical" if (degraded or low) else "ok"
        if degraded:
            degraded_count += 1
        if low:
            critical_count += 1
        robots_detail.append({
            "robot_id": r.id,
            "name": r.name,
            "color": color_to_hex(r.color),
            "battery": round(r.battery, 1),
            "battery_max": round(r.battery_max, 1),
            "degradation_pct": degradation,
            "degraded": degraded,
            "charging": r.phase == "charging",
            "phase": r.phase,
            "health": health,
            "stuck": r.stuck_alert_raised,
            "batteries_replaced": replaced.get(r.id, 0),
        })

    return {
        "degraded_threshold_pct": BATTERY_DEGRADED_WARNING_PCT,
        "totals": {
            "robots": len(robots),
            "batteries_replaced": sum(replaced.values()),
            "degraded_count": degraded_count,
            "critical_count": critical_count,
            "stuck_count": sum(1 for r in robots if r.stuck_alert_raised),
        },
        "robots": robots_detail,
        "history": history,
    }


@app.get("/api/telemetry/{robot_id}")
async def robot_telemetry(
    robot_id: int,
    bucket: str = "5 minutes",
    start: str | None = Query(None, alias="from"),
    end: str | None = Query(None, alias="to"),
    limit: int = 2000,
):
    return await db.get_telemetry(robot_id, _parse_ts(start), _parse_ts(end),
                                  bucket, min(limit, 5000))
